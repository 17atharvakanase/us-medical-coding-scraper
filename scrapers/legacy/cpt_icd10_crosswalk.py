from datetime import datetime
import ast
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.common.alert import Alert
from selenium.webdriver.common.by import By
from selenium import webdriver
import time
import threading
import pandas as pd
import json
from selenium.webdriver.common.keys import Keys
import openpyxl  # For reading/writing Excel files

import numpy as np
# import pywinauto
# from pywinauto.application import Application
# from pywinauto.keyboard import send_keys
# import pyautogui
# import pydirectinput
# pyautogui.FAILSAFE = False
import os

from selenium.webdriver.support.ui import Select
import os
import time
import  sys
import logging

# Set up the logger
logger = logging.getLogger()
logger.setLevel(logging.INFO)  # This will only log INFO, WARNING, ERROR, and CRITICAL

# Create a file handler for logging to a file
file_handler = logging.FileHandler('app.log')
file_handler.setLevel(logging.INFO)  # This ensures only INFO level logs and above are written to the file

# Create a formatter for the log entries
formatter = logging.Formatter('%(asctime)s - %(levelname)s - %(message)s', datefmt='%Y-%m-%d %H:%M:%S')
file_handler.setFormatter(formatter)

# Add the file handler to the logger
logger.addHandler(file_handler)


threadLocal = threading.local()

current_directory=os.getcwd()

ERAPATH=os.path.join(current_directory,'Config')


if not os.path.exists(ERAPATH):
    os.makedirs(ERAPATH,exist_ok=True)
    print(f'created directory {ERAPATH}')
    


def get_driver():
    driver = getattr(threadLocal, 'driver', None)
    if driver is None:
        options = webdriver.ChromeOptions()

        settings = {
            "recentDestinations": [{
                "id": "Save as PDF",
                "origin": "local",
                "account": "",
            }],
            "selectedDestinationId": "Save as PDF",
            "version": 2,
            "isHeaderFooterEnabled": False,
            "mediaSize": {
                "height_microns": 210000,
                "name": "ISO_A5",
                "width_microns": 148000,
                "custom_display_name": "A5"
            },
            "customMargins": {},
            "marginsType": 2,
            "scaling": 175,
            "scalingType": 3,
            "scalingTypePdf": 3,
            "isCssBackgroundEnabled": True
        }

        prefs = {
            "download.default_directory": ERAPATH,
            "download.prompt_for_download": False,
            "download.show_open_folder_on_save": False,  # This prevents showing the download folder on save
            "download.directory_upgrade": True,
            "plugins.always_open_pdf_externally": True,
            "printing.print_preview_sticky_settings.appState": json.dumps(settings),
            "download.show_confirm_dialog": False  # This setting prevents the confirmation dialog
        }
        options.add_experimental_option("prefs", prefs)
        options.add_argument('--window-size=1440,900')
        options.add_argument("--lang=en")
        options.add_argument("--start-maximized")
        options.add_argument("user-agent=Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/58.0.3029.110 Safari/537.3")

        options.add_experimental_option('useAutomationExtension', False)
        options.add_argument('--kiosk-printing')
        options.add_argument('--enable-print-browser')

        # options.add_argument("--headless")
        driver = webdriver.Chrome(options=options)
        
        setattr(threadLocal, 'driver', driver)
    return driver


# def check_element_loaded(driver, element_xpath, timeout=10):
#     try:
#         # Wait for the element to be present in the DOM
#         element = WebDriverWait(driver, timeout).until(
#             EC.presence_of_element_located((By.XPATH, element_xpath))
#         )
        
#         # Once the element is present, wait for it to be visible
#         element = WebDriverWait(driver, timeout).until(
#             EC.visibility_of_element_located((By.XPATH, element_xpath))
#         )
        
#         return True
#     except Exception as e:
#         print(f"Element not found or not visible. Error: {e}")
#         return False

def check_element_loaded(driver, xpath):
    try:
        element = WebDriverWait(driver, 10).until(
            EC.visibility_of_element_located((By.XPATH, xpath))
        )
        return True
    except:
        return False


def save_to_excel(data_dict, file_name):
    # Check if the file exists
    if os.path.exists(file_name):
        # Load the existing Excel file
        existing_df = pd.read_excel(file_name)
        
        # Convert the dictionary to a DataFrame
        new_df = pd.DataFrame(data_dict)
        
        # Append the new data to the existing DataFrame
        updated_df = pd.concat([existing_df, new_df], ignore_index=True)
        
        # Save the updated DataFrame back to the Excel file
        updated_df.to_excel(file_name, index=False)
        print(f"Data appended to {file_name}")
    else:
        # If the file doesn't exist, create a new one with the dictionary data
        new_df = pd.DataFrame(data_dict)
        new_df.to_excel(file_name, index=False)
        print(f"New Excel file created: {file_name}")

def read_cpt_codes(excel_file, sheet_name, column_name):
    workbook = openpyxl.load_workbook(excel_file)
    sheet = workbook[sheet_name]
    cpt_codes = []
    for row in sheet.iter_rows(min_row=2, min_col=1, max_col=sheet.max_column):
        cpt_code = row[0].value  # Assuming 'cptcode' is in the first column
        if cpt_code:
            cpt_codes.append(cpt_code)
    return cpt_codes

if __name__=='__main__':
    try:
        logging.info("Script started.")
        driver=get_driver()
        url="https://www.aapc.com/memberarea/my-purchased-items.aspx#2"
        username= "YOUR_AAPC_USERNAME"
        password="YOUR_AAPC_PASSWORD"
        driver.get(url)
        WebDriverWait(driver, 30).until(EC.element_to_be_clickable((By.XPATH,'//*[@id="signInName"]'))).click()
        WebDriverWait(driver, 30).until(EC.element_to_be_clickable((By.XPATH,'//*[@id="signInName"]'))).clear()
        WebDriverWait(driver, 30).until(EC.element_to_be_clickable((By.XPATH,'//*[@id="signInName"]'))).send_keys(str(username))
        WebDriverWait(driver, 30).until(EC.element_to_be_clickable((By.XPATH,'//*[@id="password"]'))).click()
        WebDriverWait(driver, 30).until(EC.element_to_be_clickable((By.XPATH,'//*[@id="password"]'))).clear()
        WebDriverWait(driver, 30).until(EC.element_to_be_clickable((By.XPATH,'//*[@id="password"]'))).send_keys(str(password))  
        WebDriverWait(driver, 30).until(EC.element_to_be_clickable((By.XPATH,'//button[@id="continue"]'))).click()
        try:
            driver.maximize_window()
        except:
            pass

        WebDriverWait(driver, 30).until(EC.element_to_be_clickable((By.XPATH,'//li[@id="pretoplogin"]//a[@id="linkSignOut"]'))).click()


    
        WebDriverWait(driver, 30).until(
            EC.element_to_be_clickable((By.XPATH, '//li[@id="pretoplogin"]//a[@id="linkSignOut"]'))
        ).click()
        try:
            dropdown_button = WebDriverWait(driver, 30).until(
                EC.element_to_be_clickable((By.XPATH, '//li[@id="pretoplogin"]//a[@id="linkSignOut"]'))
            )

            # Scroll the element into view and click it using JavaScript
            driver.execute_script("arguments[0].scrollIntoView(true);", dropdown_button)  # Ensure the element is in view
            driver.execute_script("arguments[0].click();", dropdown_button)  # Click using JavaScript

            # Optional: Wait for the dropdown to appear and ensure it's interactable
            dropdown_element = WebDriverWait(driver, 10).until(
                EC.visibility_of_element_located((By.XPATH, '//*[@id="MyAAPCMenuItemTop"]/div/ul'))
            )

            if dropdown_element.is_displayed():
                print("Dropdown is visible.")
                codify_link = WebDriverWait(driver, 10).until(
                    EC.element_to_be_clickable((By.XPATH, '//*[@id="lnkUserLoggedInCodifySubscription"]/a'))
                )
                codify_link.click()

        except Exception as e:
            try:
                # Step 1: Wait for and click on the "User8 Gebbs" link (which triggers the dropdown)
                user_link = WebDriverWait(driver, 10).until(
                    EC.element_to_be_clickable((By.XPATH, '//li[@id="pretoplogin"]//a[@id="linkSignOut"]'))
                )
                user_link.click()  # Click to open the dropdown

                # Step 2: Wait for the dropdown to be visible
                dropdown_menu = WebDriverWait(driver, 10).until(
                    EC.visibility_of_element_located((By.XPATH, '//*[@id="MyAAPCMenuItemTop"]/div/ul'))
                )
                
                # Step 3: Now click on the "Codify" link
                codify_link = WebDriverWait(driver, 10).until(
                    EC.element_to_be_clickable((By.XPATH, '//*[@id="lnkUserLoggedInCodifySubscription"]/a'))
                )
                codify_link.click()  # Click on the "Codify" link

                print("Codify link clicked.")

            except Exception as e:
                print(f"An error occurred: {e}")
                

        # cpt_codes = read_cpt_codes(excel_file, sheet_name, 'cptcode')
        try:
            # df = pd.read_excel(r'cpt_excel1.xls', engine='xlrd')
            df = pd.read_excel(r'cpt_not_macthing.xlsx')
        except:
            df = pd.read_excel(r'cpt_not_macthing.xlsx', engine='xlrd')

        try:
            cpt_codes = df['cptcode'].dropna().tolist()  
        except:
            cpt_codes = df['cpt'].dropna().tolist()  

        
        for cpts_count,cpts in enumerate(cpt_codes):
            cpts = str(cpts).strip()
            if len(cpts)<5:
                cpts = cpts.zfill(5)  
                          
            logging.info(f"number of counts {cpts_count} . cpts name {cpts}")

            try:
                window_after = driver.window_handles
                driver.switch_to.default_content()
                driver.switch_to.window(window_after[0])
                try:
                    search_input = WebDriverWait(driver, 2).until(EC.element_to_be_clickable((By.XPATH, '//input[@id="top_search"]')))

                except:
                    try:
                        window_after = driver.window_handles
                        driver.switch_to.default_content()
                        driver.switch_to.window(window_after[0])
                        close_button = WebDriverWait(driver, 2).until(
                            EC.element_to_be_clickable((By.XPATH, '//*[@id="autoFavoriteModal"]/div/div/div[2]/div/div[2]/button'))
                        )

                        # Click the button to close the modal
                        close_button.click()
                        search_input = WebDriverWait(driver, 2).until(EC.element_to_be_clickable((By.XPATH, '//input[@id="top_search"]')))
                    except:
                        search_input = WebDriverWait(driver, 2).until(EC.element_to_be_clickable((By.XPATH, '//*[@id="googleInput1"]')))


                    
 
                search_input.click()
                search_input.clear()
                search_input.send_keys(str(cpts.strip())) 
                search_input.send_keys(Keys.ENTER)

                try:
                    page_element = WebDriverWait(driver, 1).until(
                        EC.presence_of_element_located((By.XPATH, '//body/div[5]'))  # Or '/html/body/div[5]' based on your page structure
                    )
                    
                    # Check if the page element is visible or interactable
                    if page_element.is_displayed():
                        print("Next page has loaded.")
                    else:
                        print("Next page did not load correctly.")
                except Exception as e:
                    print(f"An error occurred while waiting for the next page to load: {e}")

        
                element_xpath='//h1[@class="pull-left"]'
                # element_xpaths=['//div[@id="collapse_cdesc"]//div[@class="tab-pane"]/text()[1]','//div[@id="collapse_desc_info"]//div[@class="tab-pane"]/text()[1]']
                element_xpaths_dict= {"HCPCS":'//div[@id="collapse_cdesc"]//div[@class="tab-pane"]','CPT':'//div[@id="collapse_desc_info"]//div[@class="tab-pane"]'}
                extracted_text=''
                extracted_text_cpt=''
                extracted_text_cpt_dict={}
                time.sleep(2)
                
                if check_element_loaded(driver, element_xpath):
                    extracted_text =driver.find_element(By.XPATH, element_xpath).text.strip()
                else:
                    try:
                        window_after = driver.window_handles
                        driver.switch_to.default_content()
                        driver.switch_to.window(window_after[0])
                        close_button = WebDriverWait(driver, 2).until(
                            EC.element_to_be_clickable((By.XPATH, '//*[@id="autoFavoriteModal"]/div/div/div[2]/div/div[2]/button'))
                        )

                        # Click the button to close the modal
                        close_button.click()
                        WebDriverWait(driver, 1).until(EC.element_to_be_clickable((By.XPATH, element_xpath)))
                        extracted_text = driver.find_element(By.XPATH, element_xpath).text.strip()
                        print(f"Extracted text after waiting: {extracted_text}")
                    
                    except:
                        pass
                    
                    
                    

                for key_,values_ in element_xpaths_dict.items():
                    if  key_ in extracted_text:  # Optional change to `and`
                        # extracted_text_cpt= WebDriverWait(driver, 2).until(EC.element_to_be_clickable((By.XPATH,values_))).get_attribute('value')
                        # extracted_text_cpt = WebDriverWait(driver, 2).until(EC.element_to_be_clickable((By.XPATH, values_))).get_attribute('value')
                        element = WebDriverWait(driver, 10).until(EC.visibility_of_element_located((By.XPATH, values_)))
                        extracted_text_cpt = element.get_attribute('value')  # For input fields

                        # If the element is not an input field, use `.text` instead of `.get_attribute('value')`
                        if not extracted_text_cpt:
                            extracted_text_cpt = element.text
                        

                        extracted_text_cpt_dict = {
                            "cpt": [cpts],               # Wrap scalar in list
                            "text": [extracted_text_cpt] # Wrap scalar in list
                        }


                        break
                if extracted_text_cpt!="" :
                    file_name = 'cpt_extracted_data1.xlsx'
                    save_to_excel(extracted_text_cpt_dict, file_name)
                else:
                    file_name = 'CPT_failed_data1.xlsx'

                    extracted_text_cpt_dict = {
                        "cpt": [cpts],               # Wrap scalar in list
                        "text": [extracted_text_cpt] # Wrap scalar in list
                    }
                    save_to_excel(extracted_text_cpt_dict, file_name)
                    
                    
            except Exception as e:
                file_name = 'CPT_failed_data1.xlsx'

                extracted_text_cpt_dict = {
                    "cpt": [cpts],               # Wrap scalar in list
                    "text": [extracted_text_cpt] # Wrap scalar in list
                }
                save_to_excel(extracted_text_cpt_dict, file_name)
                try:
                    # Switch to the alert popup
                    alert = Alert(driver)
                    
                    # Click 'No' (dismiss the alert)
                    alert.dismiss()
                except:
                    pass
                try:
                    window_after = driver.window_handles
                    driver.switch_to.default_content()
                    driver.switch_to.window(window_after[0])
                    close_button = WebDriverWait(driver, 1).until(
                        EC.element_to_be_clickable((By.XPATH, '//*[@id="autoFavoriteModal"]/div/div/div[2]/div/div[2]/button'))
                    )

                    # Click the button to close the modal
                    close_button.click()
                except:
                    pass
                # window_after = driver.window_handles
                # driver.switch_to.default_content()
                # driver.switch_to.window(window_after[0])
                # driver.refresh()
                # time.sleep(2)
                starting_window1 = driver.window_handles
                driver.switch_to.default_content()
                driver.switch_to.window(starting_window1[0])
                exc_type, exc_obj, exc_tb = sys.exc_info()
                message = "create_connection line : "+str(exc_tb.tb_lineno)+" " +str(e) 
                logging.error(message)    
             
    
 

        logging.info("Script Finished.")

        driver.quit()
    except Exception as e:
        exc_type, exc_obj, exc_tb = sys.exc_info()
        message = "create_connection line : "+str(exc_tb.tb_lineno)+" " +str(e) 
        logging.error(message)     
